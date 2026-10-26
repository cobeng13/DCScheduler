import io

import openpyxl
from openpyxl.worksheet.cell_range import CellRange
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import room_schedule
from app.main import app

URL = '/api/reports/room-schedule.xlsx'
ROOMS = [{'id': 1, 'name': 'R1'}, {'id': 2, 'name': 'R2'}]


def entry(**changes):
    return {'id': 1, 'room_id': 1, 'Course Code': 'C101', 'Section': 'A', 'Program': 'P1',
            'Faculty': 'F1', 'Room': 'R1', 'Days': 'M,W', 'Time (24 Hrs)': '08:00-09:30', **changes}


def build(entries, rooms=ROOMS, **options):
    return openpyxl.load_workbook(io.BytesIO(room_schedule.build_workbook(rooms, entries, '2026 - 2027', **options)))


def download(client, **query):
    response = client.get(URL, params={'school_year': '2026 - 2027', **query})
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return openpyxl.load_workbook(io.BytesIO(response.content))


def values(sheet):
    return '\n'.join(str(cell.value) for row in sheet for cell in row if cell.value is not None)


def test_overview_pairs_days_in_ninety_minute_bands_and_keeps_exact_times():
    book = build([entry()])
    assert book.sheetnames == ['ROOM']
    mon = book['ROOM']
    assert [cell.value for cell in mon[6]] == ['Time', 'MON\nR1', 'MON\nR2', 'Time', 'WED\nR1', 'WED\nR2']
    assert mon.max_row == 37
    assert mon['A7'].value == '7:00 AM–8:30 AM' and mon['A15'].value == '7:00 PM–8:30 PM'
    assert 'B7:B8' in set(map(str, mon.merged_cells.ranges))
    assert mon['B7'].value == 'C101\nA (P1)\n8:00 AM–9:30 AM\nF1'
    assert mon['E7'].value == mon['B7'].value
    assert [mon.cell(row, 2).value for row in (17, 28)] == ['TUE\nR1', 'FRI\nR1']
    assert [mon.cell(row, 5).value for row in (17, 28)] == ['THU\nR1', 'SAT\nR1']
    assert not any('C101' in str(cell.value) for row in mon.iter_rows(min_row=17) for cell in row)
    assert mon.freeze_panes == 'B7' and mon.print_title_cols == '$A:$A'
    assert mon.page_setup.orientation == 'landscape' and mon.page_setup.scale == 80


def test_shared_rooms_include_foreign_programs_and_selected_room_ids(clients, entry_payload):
    assert clients['alpha'].post('/api/schedule', json=entry_payload).status_code == 200
    assert clients['beta'].post('/api/schedule', json={**entry_payload, 'Program': 'P2', 'Course Code': 'FOREIGN',
        'Faculty': 'F2', 'Time (LPU Std)': '10:00a-11:00a'}).status_code == 200
    sheet = download(clients['alpha'], room_ids=[1])['ROOM']
    assert sheet.max_column == 4
    assert 'A (P1)' in values(sheet) and 'A (P2)' in values(sheet) and 'FOREIGN' in values(sheet)
    assert 'FOREIGN' not in values(download(clients['alpha'], room_ids=[2])['ROOM'])
    # Export is a view operation, not restricted to a program's assigned editor.
    assert 'C101' in values(download(clients['beta'])['ROOM'])


def test_section_export_option_adds_complete_room_schedules(clients, entry_payload):
    assert clients['alpha'].post('/api/schedule', json=entry_payload).status_code == 200
    assert clients['beta'].post('/api/schedule', json={**entry_payload, 'Program': 'P2', 'Course Code': 'P2ONLY',
        'Days': 'T', 'Faculty': 'F2'}).status_code == 200
    for scope in ['program', 'sections', 'all']:
        query = {'scope': scope, 'school_year': '2026 - 2027', 'include_rooms': True}
        if scope != 'all':
            query['program_id'] = 1
        if scope == 'sections':
            query['section_ids'] = [1]
        response = clients['alpha'].get('/api/reports/section-schedule.xlsx', params=query)
        assert response.status_code == 200, response.text if response.status_code != 200 else ''
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert book.sheetnames[:1] == ['P1']
        assert 'P2ONLY' in values(book['ROOM'])
        assert ('P2' in book.sheetnames) == (scope == 'all')


def test_quarter_hour_meetings_are_preserved_without_false_conflict():
    book = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '07:00-07:15'}),
                  entry(id=2, **{'Days': 'M', 'Course Code': 'NEXT', 'Time (24 Hrs)': '07:15-07:30'})])
    cell = book['ROOM']['B7']
    assert 'C101' in cell.value and 'NEXT' in cell.value and '7:15 AM' in cell.value
    assert cell.fill.fgColor.rgb == '00DCE6F1'
    assert 'Room TBA' not in book.sheetnames


def test_shared_band_stacks_back_to_back_classes_without_merging_or_conflict():
    sheet = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '10:30-12:00'}),
                   entry(id=2, **{'Days': 'M', 'Course Code': 'NEXT', 'Time (24 Hrs)': '12:00-13:30'})])['ROOM']
    assert 'C101' in sheet['B9'].value and 'NEXT' not in sheet['B9'].value
    assert 'C101' in sheet['B10'].value and 'NEXT' in sheet['B10'].value
    assert '10:30 AM–12:00 PM' in sheet['B10'].value
    assert '12:00 PM–1:30 PM' in sheet['B10'].value
    assert 'NEXT' in sheet['B11'].value and 'C101' not in sheet['B11'].value
    assert all(sheet[f'B{row}'].fill.fgColor.rgb == '00DCE6F1' for row in (9, 10, 11))
    assert not any(merged.min_col == 2 and merged.min_row <= 11 and merged.max_row >= 9
                   for merged in sheet.merged_cells.ranges)


def test_irregular_long_meeting_merges_touched_bands_but_preserves_exact_times():
    sheet = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '09:00-12:00'})])['ROOM']
    assert 'B8:B10' in set(map(str, sheet.merged_cells.ranges))
    assert '9:00 AM–12:00 PM' in sheet['B8'].value
    assert sheet['B7'].value is None and sheet['B11'].value is None


def test_empty_bands_between_meetings_stay_empty_and_order_is_chronological():
    entries = [entry(id=2, **{'Days': 'M', 'Course Code': 'LATER', 'Time (24 Hrs)': '13:00-14:30'}),
               entry(**{'Days': 'M', 'Time (24 Hrs)': '07:00-08:30'})]
    sheet = build(entries)['ROOM']
    assert 'C101' in sheet['B7'].value and 'LATER' in sheet['B11'].value
    assert all(sheet.cell(row, 2).value is None for row in (8, 9, 10))
    stacked = build([entry(id=2, **{'Days': 'M', 'Course Code': 'LATER', 'Time (24 Hrs)': '08:00-08:30'}),
                     entry(**{'Days': 'M', 'Time (24 Hrs)': '07:00-08:00'})])['ROOM']['B7'].value
    assert stacked.index('C101') < stacked.index('LATER')


def test_print_areas_keep_days_and_room_chunks_separate():
    rooms = [{'id': i, 'name': f'R{i}'} for i in range(1, 15)]
    sheet = build([], rooms=rooms)['ROOM']
    assert sheet['B6'].value == 'MON\nR1' and sheet['C6'].value == 'MON\nR2'
    assert sheet['O6'].value == 'MON\nR14'  # Natural room order, not R1/R10/R11.
    assert sheet['Q6'].value == 'WED\nR1'
    areas = [CellRange(area.split('!')[-1]) for area in str(sheet.print_area).split(',')]
    assert len(areas) == 18  # Three room chunks for each of six days.
    assert all(area.max_col - area.min_col + 1 <= 6 for area in areas)
    assert all(area.max_row - area.min_row == 9 for area in areas)
    assert sheet.print_title_cols == '$A:$A'
    assert sheet.page_setup.fitToWidth is None and sheet.page_setup.scale == 80


def test_midnight_clips_bands_without_shifting_the_seven_am_anchor():
    sheet = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '00:00-01:15'}),
                   entry(id=2, **{'Days': 'M', 'Course Code': 'LATE', 'Time (24 Hrs)': '23:30-24:00'})])['ROOM']
    assert sheet['A7'].value == '12:00 AM–1:00 AM'
    assert sheet['A8'].value == '1:00 AM–2:30 AM'
    assert sheet['B7'].value == 'C101\nA (P1)\n12:00 AM–1:15 AM\nF1'
    assert sheet.cell(sheet.max_row, 1).value == '11:30 PM–12:00 AM'
    assert 'LATE' in values(sheet)
    assert '7:00 AM–8:30 AM' in values(sheet)


def test_dense_labels_fail_explicitly_instead_of_clipping_or_exceeding_excel_limits():
    with pytest.raises(HTTPException) as error:
        build([entry(**{'Days': 'M', 'Faculty': 'Long faculty name ' * 800})])
    assert error.value.status_code == 413
    assert 'fit in Excel' in error.value.detail


def test_true_overlaps_are_visible_and_neither_booking_is_lost():
    book = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '07:00-08:00'}),
                  entry(id=2, **{'Days': 'M', 'Course Code': 'CONFLICT', 'Time (24 Hrs)': '07:30-08:30'})])
    mon = book['ROOM']
    assert 'C101' in mon['B7'].value
    assert 'CONFLICT' in mon['B7'].value
    assert mon['B7'].fill.fgColor.rgb == '00FCE4D6'
    assert mon['B8'].value is None  # An 8:30 end does not occupy the next band.


def test_boundaries_extend_grid_and_sunday_is_never_dropped():
    book = build([entry(**{'Days': 'Su', 'Time (24 Hrs)': '06:15-06:45'}),
                  entry(id=2, **{'Days': 'M', 'Time (24 Hrs)': '21:15-22:00'})])
    sheet = book['ROOM']
    assert 'SUN\nR1' in values(sheet)
    assert sheet['A7'].value == '5:30 AM–7:00 AM'
    assert '6:15 AM–6:45 AM' in values(book['ROOM'])
    assert sheet.cell(sheet.max_row, 1).value == '8:30 PM–10:00 PM'
    assert 'SUN\nR1' in values(build([], include_sunday=True)['ROOM'])
    assert 'SUN\nR1' not in values(build([])['ROOM'])


def test_tba_invalid_or_unavailable_assignments_are_listed_separately():
    book = build([entry(**{'room_id': None, 'Room': 'TBA'}),
                  entry(id=2, **{'Course Code': 'NO TIME', 'Time (24 Hrs)': 'TBA'}),
                  entry(id=3, **{'Course Code': 'NO DAY', 'Days': 'TBA'}),
                  entry(id=4, **{'Course Code': 'INVALID', 'Time (24 Hrs)': '25:00-26:00'})])
    assert book['Room TBA'].max_row == 6
    text = values(book['Room TBA'])
    assert all(word in text for word in ['C101', 'NO TIME', 'NO DAY', 'INVALID'])
    assert 'C101' not in values(book['ROOM'])


def test_empty_room_catalog_safe_text_and_sheet_name_collisions():
    assert build([], rooms=[])['ROOM']['B7'].value == 'No rooms to export.'
    book = build([entry(**{'Course Code': '=HYPERLINK("bad")'})], institution='=1+1')
    assert book['ROOM']['A1'].data_type == 's' and book['ROOM']['B7'].data_type == 's'
    book.create_sheet('Room TBA')
    room_schedule.append_sheets(book, ROOMS, [entry(**{'Days': 'TBA'})], '2026')
    assert 'ROOM (2)' in book.sheetnames and 'Room TBA (2)' in book.sheetnames


@pytest.mark.parametrize('query', [{'room_ids': [999]}, {'room_ids': [1, 999]}, {'school_year': ' '}, {'title': ' '}])
def test_invalid_export_requests_are_rejected(clients, query):
    assert clients['alpha'].get(URL, params={'school_year': '2026', **query}).status_code == 422


def test_export_requires_login_and_is_read_only(clients, entry_payload):
    with TestClient(app) as anonymous:
        assert anonymous.get(URL, params={'school_year': '2026'}).status_code == 401
    alpha = clients['alpha']
    created = alpha.post('/api/schedule', json=entry_payload).json()
    activity = alpha.get('/api/activity').json()
    download(alpha, institution='Custom institution', college='Custom college', title='Custom rooms')
    assert alpha.get('/api/schedule').json() == [created]
    assert alpha.get('/api/activity').json() == activity


def test_export_limits_prevent_excessively_large_grids():
    with pytest.raises(HTTPException) as error:
        build([], rooms=[{'id': i, 'name': str(i)} for i in range(5000)])
    assert error.value.status_code == 413


def test_lpu_fallback_invalid_clock_minutes_and_midnight_are_handled():
    book = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '', 'Time (LPU Std)': '11:30p-12:00a'}),
                  entry(id=2, **{'Course Code': 'BAD MINUTES', 'Time (24 Hrs)': '08:99-10:00', 'Time (LPU Std)': 'TBA'})])
    # LPU's same-day format cannot express an overnight interval, so use an
    # authoritative 24-hour value ending at midnight for a valid boundary.
    assert 'BAD MINUTES' in values(book['Room TBA'])
    book = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '23:30-24:00'})])
    assert '11:30 PM–12:00 AM' in values(book['ROOM'])
    book = build([entry(**{'Days': 'M', 'Time (24 Hrs)': '', 'Time (LPU Std)': '8:15a-9:00a'})])
    assert '8:15 AM–9:00 AM' in values(book['ROOM'])


def test_ignored_rooms_and_their_bookings_are_excluded_from_both_export_paths(clients, entry_payload):
    alpha = clients['alpha']
    assert alpha.post('/api/schedule', json={**entry_payload, 'Course Code': 'EXCLUDED'}).status_code == 200
    assert alpha.post('/api/schedule', json={**entry_payload, 'Course Code': 'KEPT', 'Room': 'R2', 'Days': 'T'}).status_code == 200
    assert clients['admin'].put('/api/admin/rules', json={'version': 1, 'ignoreRoom': False,
        'ignoreFaculty': False, 'ignoreRoomIds': [1]}).status_code == 200
    # Explicit old selections cannot bypass the server's exclusion either.
    for query in [{}, {'room_ids': [1, 2]}, {'room_ids': [1]}]:
        book = download(alpha, **query)
        text = '\n'.join(values(sheet) for sheet in book)
        assert 'R1' not in text and 'EXCLUDED' not in text
        assert 'Room TBA' not in book.sheetnames
        if query.get('room_ids') != [1]:
            assert 'KEPT' in text and book['ROOM']['B6'].value == 'MON\nR2'
    response = alpha.get('/api/reports/section-schedule.xlsx', params={
        'scope': 'program', 'program_id': 1, 'school_year': '2026', 'include_rooms': True})
    assert response.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert 'EXCLUDED' in values(book['P1'])  # Section tables still describe their classes.
    assert 'EXCLUDED' not in '\n'.join(values(sheet) for sheet in book if sheet.title != 'P1')
    assert book['ROOM']['B6'].value == 'MON\nR2'


def test_global_ignore_switch_does_not_exclude_untagged_rooms(clients, entry_payload):
    assert clients['alpha'].post('/api/schedule', json=entry_payload).status_code == 200
    assert clients['admin'].put('/api/admin/rules', json={'version': 1, 'ignoreRoom': True,
        'ignoreFaculty': False, 'ignoreRoomIds': []}).status_code == 200
    assert 'C101' in values(download(clients['alpha'])['ROOM'])
