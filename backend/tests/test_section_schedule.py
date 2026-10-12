import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import models, section_schedule
from app.db import SessionLocal
from app.main import app

URL = '/api/reports/section-schedule.xlsx'


def state(courses=None):
    return {'selectedTerm': 'First Semester', 'sectionYearLevels': {'a': 'First Year'},
            'sectionCurriculumIds': {'a': '2026'}, 'yearLevelCurriculumIds': {},
            'curricula': [{'id': '2026', 'name': 'Curriculum 2026 - 2027', 'courses': courses or [
                {'courseCode': 'C101', 'courseDescription': 'Lecture', 'totalUnits': 3, 'hours': 3, 'yearLevel': 'First Year', 'semester': 'First Semester'},
                {'courseCode': 'C102 Lab', 'courseDescription': 'Laboratory', 'totalUnits': 2, 'hours': 6, 'yearLevel': 'First Year', 'semester': 'First Semester'},
                {'courseCode': 'OTHER-YEAR', 'courseDescription': 'Excluded', 'totalUnits': 5, 'hours': 5, 'yearLevel': 'Second Year', 'semester': 'First Semester'},
                {'courseCode': 'OTHER-TERM', 'courseDescription': 'Excluded', 'totalUnits': 5, 'hours': 5, 'yearLevel': 'First Year', 'semester': 'Second Semester'},
            ]}]}


def configure(program_id=1, curriculum=None):
    with SessionLocal() as db:
        db.get(models.Program, program_id).settings_json = json.dumps({'curriculumState': curriculum or state()})
        db.commit()


def export(client, scope='program', **query):
    params = {'scope': scope, 'school_year': '2026 - 2027', **query}
    if scope != 'all':
        params.setdefault('program_id', 1)
    return client.get(URL, params=params)


def workbook(response):
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_curriculum_complete_export_merges_meetings_and_totals_once(clients, entry_payload):
    configure()
    alpha = clients['alpha']
    assert alpha.post('/api/schedule', json={**entry_payload, 'Days': 'M,W'}).status_code == 200
    assert alpha.post('/api/schedule', json={**entry_payload, 'Days': 'Th', 'Time (LPU Std)': '10:00a-11:00a', 'Room': 'R2', 'Faculty': 'F2'}).status_code == 200
    sheet = workbook(export(alpha))['P1']
    assert sheet['A8'].value == 'A' and sheet['B8'].value == 'Curriculum 2026 - 2027'
    assert [cell.value for cell in sheet[9]] == section_schedule.HEADERS
    assert [sheet.cell(10, column).value for column in range(1, 9)] == ['C101', 'Lecture', 3, 3, '08:00 - 09:00 LEC', 'MON/WED', 'R1', 'F1']
    assert [sheet.cell(11, column).value for column in range(5, 9)] == ['10:00 - 11:00 LEC', 'THU', 'R2', 'F2']
    assert {'A10:A11', 'B10:B11', 'C10:C11', 'D10:D11'}.issubset(set(map(str, sheet.merged_cells.ranges)))
    assert 'H10:H11' not in set(map(str, sheet.merged_cells.ranges))
    assert [sheet.cell(12, column).value for column in range(1, 9)] == ['C102 Lab', 'Laboratory', 2, 6, 'TBA', 'TBA', 'TBA', 'TBA']
    assert sheet['C13'].value == '=SUM(C10:C12)' and sheet['D13'].value == '=SUM(D10:D12)'
    assert sum(sheet.cell(row, 3).value or 0 for row in range(10, 13)) == 5
    assert sum(sheet.cell(row, 4).value or 0 for row in range(10, 13)) == 9
    assert sheet.page_setup.orientation == 'landscape' and sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == '$1:$6' and sheet.print_area == "'P1'!$A$1:$H$13"
    assert not any('Online' in str(cell.value) or 'F2F' in str(cell.value) for row in sheet for cell in row)


def test_all_programs_have_separate_sheets_and_isolated_same_named_sections(clients, entry_payload):
    configure(1)
    other = state()
    other['curricula'][0]['name'] = 'Other curriculum'
    other['curricula'][0]['courses'][0]['courseCode'] = 'OTHER-PROGRAM'
    configure(2, other)
    assert clients['alpha'].post('/api/schedule', json=entry_payload).status_code == 200
    result = workbook(export(clients['alpha'], 'all'))
    assert result.sheetnames == ['P1', 'P2']
    assert result['P1']['A10'].value == 'C101'
    assert result['P2']['A10'].value == 'OTHER-PROGRAM'
    assert result['P2']['E10'].value == 'TBA' and result['P2']['H10'].value == 'TBA'
    # Read-only access is the same as other Online schedule exports.
    assert export(clients['beta'], program_id=1).status_code == 200


def test_selected_sections_scope_and_each_section_has_own_totals(clients):
    configure()
    alpha = clients['alpha']
    response = alpha.post('/api/sections', json={'program_id': 1, 'name': 'B'})
    assert response.status_code == 200
    section_id = response.json()['id']
    with SessionLocal() as db:
        program = db.get(models.Program, 1)
        settings = json.loads(program.settings_json)
        settings['curriculumState']['sectionCurriculumIds']['b'] = '2026'
        settings['curriculumState']['sectionYearLevels']['b'] = 'First Year'
        program.settings_json = json.dumps(settings)
        db.commit()
    sheet = workbook(export(alpha, 'sections', section_ids=[section_id]))['P1']
    assert sheet['A8'].value == 'B'
    full = workbook(export(alpha))['P1']
    totals = [cell for row in full for cell in row if cell.data_type == 'f']
    assert len(totals) == 4 and len(full.row_breaks.brk) == 1
    assert totals[2].value != totals[0].value


@pytest.mark.parametrize('scope,query', [
    ('sections', {}), ('sections', {'section_ids': [2]}), ('sections', {'section_ids': [999]}),
    ('program', {'section_ids': [1]}), ('all', {'program_id': 1}), ('all', {'section_ids': [1]}),
    ('other', {}), ('program', {'program_id': 999}),
])
def test_invalid_filters_are_rejected(clients, scope, query):
    response = export(clients['alpha'], scope, **query)
    assert response.status_code in (404, 422), response.text


def test_export_requires_login_and_does_not_change_data(clients, entry_payload):
    with TestClient(app) as anonymous:
        assert export(anonymous, 'all').status_code == 401
    alpha = clients['alpha']
    entry = alpha.post('/api/schedule', json=entry_payload).json()
    events = alpha.get('/api/activity').json()
    workbook(export(alpha))
    assert alpha.get('/api/schedule').json() == [entry]
    assert alpha.get('/api/activity').json() == events


def test_missing_assignments_are_explicit_and_entered_extras_are_preserved(clients, entry_payload):
    alpha = clients['alpha']
    assert alpha.post('/api/schedule', json=entry_payload).status_code == 200
    sheet = workbook(export(alpha))['P1']
    assert 'entered classes only' in sheet['B8'].value
    assert sheet['A10'].value == 'C101'
    configure()
    assert alpha.post('/api/schedule', json={**entry_payload, 'Course Code': 'EXTRA', 'Days': 'T'}).status_code == 200
    sheet = workbook(export(alpha))['P1']
    assert [sheet.cell(row, 1).value for row in range(10, 12)] == ['C101', 'C102 Lab']
    assert 'Additional entered classes' in sheet['A12'].value and sheet['A13'].value == 'EXTRA'


def test_sheet_names_are_safe_unique_and_truncated():
    used = set()
    names = [section_schedule.sheet_name(name, used) for name in ['Bad/Name', 'Bad:Name', 'bad/name', 'X' * 50, 'X' * 51, "''"]]
    assert len({name.lower() for name in names}) == len(names)
    assert all(len(name) <= 31 and name for name in names)
    assert not any('/' in name or ':' in name for name in names)


def test_text_is_literal_and_headings_are_customizable(clients):
    courses = [{'courseCode': '=1+1', 'courseDescription': '=HYPERLINK("bad")', 'totalUnits': 1.5, 'hours': 4.5,
                'yearLevel': 'First Year', 'semester': 'First Semester'}]
    configure(curriculum=state(courses))
    sheet = workbook(export(clients['alpha'], institution='Custom University', college='Custom College', title='Custom Schedule'))['P1']
    assert sheet['A1'].value == 'Custom University' and sheet['A3'].value == 'Custom College' and sheet['A4'].value == 'Custom Schedule'
    assert sheet['A10'].value == '=1+1' and sheet['A10'].data_type == 's'
    assert sheet['B10'].data_type == 's' and sheet['C10'].value == 1.5 and sheet['D10'].value == 4.5


def test_multiple_curricula_require_assignment_and_empty_program_is_exported(clients):
    current = state()
    current['sectionCurriculumIds'] = {}
    current['curricula'].append({**current['curricula'][0], 'id': 'old', 'name': 'Old curriculum'})
    configure(curriculum=current)
    sheet = workbook(export(clients['alpha']))['P1']
    assert 'Curriculum not assigned' in sheet['B8'].value and sheet['C11'].value == 'TBA'
    with SessionLocal() as db:
        db.add(models.Program(name='Empty program'))
        db.commit()
    sheet = workbook(export(clients['alpha'], 'all'))['Empty program']
    assert sheet['A8'].value == 'No sections in this program.'


def test_year_level_default_selects_correct_curriculum_and_avoids_other_years():
    current = state()
    current['sectionCurriculumIds'] = {}
    current['yearLevelCurriculumIds'] = {'first year': '2026'}
    courses, label, complete = section_schedule.assigned_courses(current, 'A')
    assert complete and [course['courseCode'] for course in courses] == ['C101', 'C102 Lab']
    current['sectionYearLevels'] = {}
    assert section_schedule.assigned_courses(current, 'A')[2] is False


def test_different_meetings_same_faculty_merge_only_shared_columns(clients, entry_payload):
    configure()
    alpha = clients['alpha']
    assert alpha.post('/api/schedule', json=entry_payload).status_code == 200
    assert alpha.post('/api/schedule', json={**entry_payload, 'Days': 'F', 'Time (LPU Std)': '11:00a-12:00p', 'Room': 'R2'}).status_code == 200
    sheet = workbook(export(alpha))['P1']
    assert 'H10:H11' in set(map(str, sheet.merged_cells.ranges))
    assert sheet['G10'].value == 'R1' and sheet['G11'].value == 'R2'


def test_laboratory_in_subject_name_does_not_imply_lab_component():
    entry = {'Time (24 Hrs)': '08:00-09:00', 'Days': 'M', 'Room': 'R1', 'Faculty': 'F1'}
    meetings = section_schedule.meeting_rows([entry], {'code': 'MLS101', 'description': 'Introduction to Medical Laboratory Science'})
    assert meetings[0]['time'] == '08:00 - 09:00'


def test_nonfinite_or_conflicting_course_values_are_rejected(clients):
    courses = state()['curricula'][0]['courses']
    courses[0]['hours'] = float('inf')
    configure(curriculum=state(courses))
    assert export(clients['alpha']).status_code == 422
    courses[0]['hours'] = 3
    courses.append({**courses[0], 'totalUnits': 4})
    configure(curriculum=state(courses))
    assert export(clients['alpha']).status_code == 422
