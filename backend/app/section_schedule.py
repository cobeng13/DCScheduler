"""Formatted, curriculum-complete Excel schedules for Online programs."""
import io
import json
import math
import re
from collections import defaultdict

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from . import time_utils

HEADERS = ['Course Code', 'Course Description', 'Units', '# of Hours', 'Time', 'Days', 'Room', 'Faculty']
DEFAULT_INSTITUTION = 'LYCEUM OF THE PHILIPPINES UNIVERSITY - BATANGAS'
DEFAULT_COLLEGE = 'COLLEGE OF ALLIED MEDICAL PROFESSIONS'
DAYS = ['M', 'T', 'W', 'Th', 'F', 'Sa', 'Su']
DAY_LABELS = dict(zip(DAYS, ['MON', 'TUE', 'WED', 'THU', 'FRI', 'SAT', 'SUN']))


def key(value):
    return str(value or '').strip().lower()


def natural(value):
    return [(1, int(part)) if part.isdigit() else (0, part.lower()) for part in re.split(r'(\d+)', value)]


def number(value, label):
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise HTTPException(422, f'{label} is missing or invalid.') from None
    if not math.isfinite(result) or result < 0:
        raise HTTPException(422, f'{label} must be a finite, nonnegative number.')
    return result


def assigned_courses(state, section):
    """Never combine versions or years when a section assignment is ambiguous."""
    year = state.get('sectionYearLevels', {}).get(key(section), '')
    identity = state.get('sectionCurriculumIds', {}).get(key(section)) or state.get('yearLevelCurriculumIds', {}).get(key(year))
    catalogs = state.get('curricula', [])
    catalog = next((item for item in catalogs if item.get('id') == identity), None)
    if not identity and len(catalogs) == 1:
        catalog = catalogs[0]
    if not catalog:
        return [], 'Curriculum not assigned (entered classes only)', False
    courses = [course for course in catalog.get('courses', []) if course.get('semester') == state.get('selectedTerm', 'First Semester')]
    years = {key(course.get('yearLevel')) for course in courses if course.get('yearLevel')}
    if not year and len(years) > 1:
        return [], 'Year level not assigned (entered classes only)', False
    if year:
        courses = [course for course in courses if key(course.get('yearLevel')) == key(year)]
    return courses, catalog.get('name') or 'Assigned curriculum', True


def section_courses(state, section, entries):
    courses, label, complete = assigned_courses(state, section)
    result = {}
    for course in courses:
        code = course.get('courseCode', '')
        if not key(code):
            raise HTTPException(422, f'{section}: an assigned curriculum course has no course code.')
        item = {'code': code, 'description': course.get('courseDescription', ''),
                'units': number(course.get('totalUnits'), f'{section} / {code} units'),
                'hours': number(course.get('hours'), f'{section} / {code} required hours'), 'meetings': []}
        if key(code) in result and result[key(code)] != item:
            raise HTTPException(422, f'{section}: conflicting curriculum values for {code}. Review the assigned curriculum.')
        result[key(code)] = item
    groups = defaultdict(list)
    for entry in entries:
        groups[key(entry['Course Code'])].append(entry)
    for code, group in sorted(groups.items(), key=lambda item: natural(item[0])):
        if code not in result:
            first = group[0]
            values = {(number(row['Units'], 'Class units'), number(row['# of Hours'], 'Class hours')) for row in group}
            if len(values) != 1:
                raise HTTPException(422, f'{section}: inconsistent units or hours for {first["Course Code"]}. Review these classes before exporting.')
            units, hours = next(iter(values))
            result[code] = {'code': first['Course Code'], 'description': first['Course Description'],
                            'units': units, 'hours': hours, 'meetings': [], 'extra': True}
        result[code]['meetings'] = meeting_rows(group, result[code])
    return list(result.values()), label, complete


def meeting_rows(entries, course):
    grouped = {}
    for entry in entries:
        try:
            start, end = time_utils.parse_time_range(entry.get('Time (24 Hrs)') or '')
        except ValueError:
            try:
                _, _, start, end = time_utils.parse_time_lpu(entry.get('Time (LPU Std)') or '')
            except ValueError:
                start = end = None
        days = time_utils.normalize_days(entry.get('Days') or '')
        if not days:
            start = end = None
        signature = (start, end, entry.get('Room') or 'TBA', entry.get('Faculty') or 'TBA')
        grouped.setdefault(signature, set()).update(days)
    result = []
    kind = re.search(r'\b(lec|lab)\b', course['code'], re.I) or re.search(r'\b(lec|lecture|lab|laboratory)\s*\)?$', course['description'], re.I)
    suffix = (' LAB' if kind.group().lower().startswith('lab') else ' LEC') if kind else ''
    clock = lambda minutes: f'{minutes // 60 % 12 or 12:02d}:{minutes % 60:02d}'
    for (start, end, room, faculty), days in grouped.items():
        ordered = [day for day in DAYS if day in days]
        result.append({'time': f'{clock(start)} - {clock(end)}{suffix}' if start is not None else 'TBA',
                       'days': '/'.join(DAY_LABELS[day] for day in ordered) or 'TBA', 'room': room, 'faculty': faculty,
                       'sort': (DAYS.index(ordered[0]) if ordered else 7, start if start is not None else 1440, room, faculty)})
    return sorted(result, key=lambda row: row['sort'])


def sheet_name(name, used):
    base = re.sub(r'[\[\]:*?/\\\x00-\x1f]', '_', name).strip(" '")[:31] or 'Program'
    candidate, index = base, 2
    while candidate.lower() in used:
        suffix = f' ({index})'
        candidate = base[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def write_value(sheet, row, column, value):
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub(' ', value)
        cell.data_type = 's'  # Untrusted labels must never become Excel formulas.
    else:
        cell.value = value
    return cell


def wrapped_height(values, widths):
    lines = max((sum(max(1, math.ceil(len(line) / max(1, width - 2))) for line in str(value or '').split('\n'))
                 for value, width in zip(values, widths)), default=1)
    return max(21, 16 * lines + 5)


def build_workbook(programs, school_year, institution=DEFAULT_INSTITUTION, college=DEFAULT_COLLEGE, title='SCHEDULE OF CLASSES', room_export=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    used = set()
    widths = [18, 42, 8, 11, 24, 14, 16, 28]
    edge = Side(style='thin', color='808080')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    total_rows = 0
    for program in sorted(programs, key=lambda program: natural(program['name'])):
        sheet = workbook.create_sheet(sheet_name(program['name'], used))
        state = json.loads(program.get('settings_json') or '{}').get('curriculumState', {})
        term = state.get('selectedTerm', 'First Semester')
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = '14'  # Match the supplied CAMP workbook.
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = sheet.page_margins.right = 0.39
        sheet.page_margins.top = sheet.page_margins.bottom = 0.2
        sheet.print_title_rows = '1:6'
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
        for row, text in [(1, institution), (3, college), (4, title), (5, f'{term.upper()} SY {school_year}'), (6, program['name'])]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            cell = write_value(sheet, row, 1, text)
            cell.font = Font(name='Arial Narrow', size=12, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            sheet.row_dimensions[row].height = max(19, 16 * math.ceil(len(text) / 120))
        sheet.row_dimensions[2].height = 8
        row = 8
        sections = sorted(program['sections'], key=lambda section: natural(section['name']))
        if not sections:
            sheet.merge_cells('A8:H8')
            write_value(sheet, 8, 1, 'No sections in this program.')
        for section_index, section in enumerate(sections):
            courses, curriculum_label, complete = section_courses(state, section['name'], section['entries'])
            write_value(sheet, row, 1, section['name']).font = Font(name='Arial Narrow', size=12, bold=True)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            cell = write_value(sheet, row, 2, curriculum_label)
            cell.font = Font(name='Arial Narrow', size=12, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            sheet.row_dimensions[row].height = 24
            row += 1
            for column, text in enumerate(HEADERS, 1):
                cell = write_value(sheet, row, column, text)
                cell.font = Font(name='Arial Narrow', size=10, bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='F2F2F2')
                cell.border = border
            sheet.row_dimensions[row].height = 27
            row += 1
            first_detail = row
            extras_labelled = False
            for course in courses:
                if complete and course.get('extra') and not extras_labelled:
                    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    cell = write_value(sheet, row, 1, 'Additional entered classes (not in the assigned curriculum)')
                    cell.font = Font(name='Arial Narrow', size=10, italic=True)
                    sheet.row_dimensions[row].height = 20
                    row += 1
                    extras_labelled = True
                first_course = row
                meetings = course['meetings'] or [{'time': 'TBA', 'days': 'TBA', 'room': 'TBA', 'faculty': 'TBA'}]
                for meeting_index, meeting in enumerate(meetings):
                    values = [course['code'], course['description'], course['units'], course['hours'],
                              meeting['time'], meeting['days'], meeting['room'], meeting['faculty']]
                    for column, value in enumerate(values, 1):
                        cell = write_value(sheet, row, column, value if meeting_index == 0 or column >= 5 else None)
                        cell.font = Font(name='Calibri', size=12)
                        cell.alignment = Alignment(horizontal='center' if column in (3, 4, 5, 6, 7) else 'left', vertical='center', wrap_text=True)
                        cell.border = border
                        if column in (3, 4):
                            cell.number_format = 'General'
                    sheet.row_dimensions[row].height = wrapped_height(values, widths)
                    row += 1
                if len(meetings) > 1:
                    for column in (1, 2, 3, 4):
                        sheet.merge_cells(start_row=first_course, start_column=column, end_row=row - 1, end_column=column)
                    if len({meeting['faculty'] for meeting in meetings}) == 1:
                        sheet.merge_cells(start_row=first_course, start_column=8, end_row=row - 1, end_column=8)
                total_rows += len(meetings)
                if total_rows > 50000:
                    raise HTTPException(413, 'This export exceeds 50,000 meeting rows. Export fewer sections or one program at a time.')
            if not courses:
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                write_value(sheet, row, 1, 'No curriculum courses for this semester.' if complete else 'Assign a curriculum and year level in Edit Sections to include the full course list.')
                sheet.row_dimensions[row].height = 30
                row += 1
            last_detail = row - 1
            for column in range(1, 9):
                cell = write_value(sheet, row, column, 'Total Units' if column == 1 else None)
                cell.font = Font(name='Arial Narrow', size=12, bold=True)
                cell.alignment = Alignment(horizontal='center' if column in (3, 4) else 'left', vertical='center')
                cell.border = border
                if column in (3, 4):
                    cell.number_format = 'General'
            for column, letter in [(3, 'C'), (4, 'D')]:
                if courses or complete:
                    sheet.cell(row, column, f'=SUM({letter}{first_detail}:{letter}{last_detail})')
                else:
                    write_value(sheet, row, column, 'TBA')
            sheet.row_dimensions[row].height = 22
            if section_index < len(sections) - 1:
                sheet.row_breaks.append(Break(id=row))
            row += 2
        sheet.print_area = f'A1:H{max(8, row - 2)}'
    if not workbook.worksheets:
        raise HTTPException(422, 'There are no programs to export.')
    if room_export is not None:
        from .room_schedule import append_sheets
        append_sheets(workbook, room_export['rooms'], room_export['entries'], school_year, institution, college)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
