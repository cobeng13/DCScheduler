"""Secretary-style room overview with 90-minute display bands and exact times."""
import io
import math
from copy import copy
from collections import defaultdict

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import time_utils
from .section_schedule import DEFAULT_COLLEGE, DEFAULT_INSTITUTION, DAYS, DAY_LABELS, natural, sheet_name, write_value, wrapped_height

BAND_MINUTES = 90
DAY_START = 420  # Anchor at 7 AM, rather than at midnight.
DAY_END = 1230  # Nine bands, ending at 8:30 PM.
MAX_CELLS = 250000
MAX_COLUMNS = 16384
ROOMS_PER_PAGE = 6


def band_start(minute):
    return max(0, DAY_START + (minute - DAY_START) // BAND_MINUTES * BAND_MINUTES)


def band_end(minute):
    return min(1440, DAY_START + math.ceil((minute - DAY_START) / BAND_MINUTES) * BAND_MINUTES)


def clock(minutes):
    hour = minutes // 60 % 12 or 12
    return f'{hour}:{minutes % 60:02d}' if minutes % 60 else str(hour)


def interval(entry):
    try:
        text = entry.get('Time (24 Hrs)') or ''
        start, end = time_utils.parse_time_range(text)
        for index, part in enumerate(text.strip().split('-')):
            hour, minute = map(int, part.split(':'))
            if minute >= 60 or hour >= 24 and not (index == 1 and hour == 24 and minute == 0):
                raise ValueError('Invalid clock time')
    except ValueError:
        try:
            _, _, start, end = time_utils.parse_time_lpu(entry.get('Time (LPU Std)') or '')
        except ValueError:
            return None
    # Legacy imported values must not extend Excel beyond a single calendar day.
    if not 0 <= start < end <= 1440:
        return None
    return start, end


def meeting_days(entry):
    value = entry.get('Days') or ''
    return set() if time_utils.is_tba(value) else time_utils.normalize_days(value)


def label(entry, start, end):
    return '\n'.join([entry.get('Course Code') or 'TBA',
                      entry.get('Section') or 'TBA',
                      f'{clock(start)}-{clock(end)}', entry.get('Faculty') or 'TBA'])


def configure_sheet(sheet, panel_width, title, school_year):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B7'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = '8'  # A3; room columns continue on additional pages.
    sheet.page_setup.scale = 80
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    # Each print area contains one day's header and at most six room columns.
    # The left time column has the same bands at the same rows as the right panel.
    sheet.print_title_cols = 'A:A'
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.65
    sheet.page_margins.bottom = 0.55
    sheet.oddHeader.center.text = f'{title[:100]}\nSY {school_year}'.replace('&', '&&')
    sheet.oddHeader.center.size = 10
    sheet.oddFooter.left.text = 'Display bands only; refer to exact meeting times. Red = actual overlap.'
    sheet.oddFooter.left.size = 9
    sheet.oddFooter.right.text = 'Page &P of &N'
    for column in range(1, panel_width * 2 + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 25


def append_sheets(workbook, rooms, entries, school_year, institution=DEFAULT_INSTITUTION,
                  college=DEFAULT_COLLEGE, title='ROOM SCHEDULE', include_sunday=False):
    rooms = sorted(rooms, key=lambda room: natural(room['name']))
    room_ids = {room['id'] for room in rooms}
    used = {sheet.title.lower() for sheet in workbook.worksheets}
    bookings, unplotted = [], []
    for entry in entries:
        times = interval(entry)
        days = meeting_days(entry)
        if entry.get('room_id') not in room_ids or not times or not days:
            unplotted.append(entry)
        else:
            bookings.append((entry, *times, days))
    bookings.sort(key=lambda booking: (booking[1], booking[2], natural(booking[0].get('Course Code') or ''),
                                       str(booking[0].get('id') or '')))
    start = min([DAY_START] + [band_start(begin) for _, begin, _, _ in bookings])
    end = max([DAY_END] + [band_end(finish) for _, _, finish, _ in bookings])
    # Midnight clips the first/last band; the remaining boundaries stay 7 AM anchored.
    boundaries = [start] + [minute for minute in range(DAY_START % BAND_MINUTES, 1441, BAND_MINUTES) if start < minute < end] + [end]
    slots = list(zip(boundaries, boundaries[1:]))
    days = DAYS if include_sunday or any('Su' in item[3] for item in bookings) else DAYS[:-1]
    panel_width = max(2, len(rooms) + 1)
    if (panel_width * 2 > MAX_COLUMNS or panel_width * len(slots) * len(days) > MAX_CELLS
            or len(entries) > 50000):
        raise HTTPException(413, 'This export is too large. Select fewer rooms.')
    edge = Side(style='thin', color='B8C4D2')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    occupied = defaultdict(list)
    preferred = {}
    occupancy_count = 0
    for index, (entry, begin, finish, booked_days) in enumerate(bookings):
        touched = [offset for offset, (left, right) in enumerate(slots) if begin < right and finish > left]
        # Reserve each meeting's largest band so compaction can never hide it.
        # max() keeps the earlier band when the occupied minutes are equal.
        preferred[index] = max(touched, key=lambda offset: min(finish, slots[offset][1]) - max(begin, slots[offset][0]))
        occupancy_count += len(booked_days) * len(touched)
        if occupancy_count > MAX_CELLS:
            raise HTTPException(413, 'This export has too many bookings. Select fewer rooms.')
        for day in booked_days:
            for offset in touched:
                occupied[day, entry['room_id'], offset].append(index)
    conflicts = {}
    for key, indexes in occupied.items():
        left, right = slots[key[2]]
        latest_end = left
        conflicts[key] = False
        for index in indexes:  # Bookings are in start-time order.
            _, begin, finish, _ = bookings[index]
            if max(left, begin) < latest_end:
                conflicts[key] = True
                break
            latest_end = min(right, finish)

    for key, indexes in occupied.items():
        if conflicts[key]:
            continue
        left, right = slots[key[2]]
        minutes = {index: min(bookings[index][2], right) - max(bookings[index][1], left) for index in indexes}
        # Omit minor leading/trailing fragments, including in otherwise empty
        # bands. Always keep a meeting's reserved band so short classes survive.
        selected = [index for index in indexes if preferred[index] == key[2]
                    or 2 * minutes[index] > right - left]
        winner = max(indexes, key=lambda index: minutes[index])
        # A shared 45/45 boundary goes to the earlier meeting. An isolated
        # exactly-half fragment is omitted unless it is the reserved band.
        if not selected and len(indexes) > 1 and 2 * minutes[winner] == right - left:
            selected = [winner]
        occupied[key] = selected

    sheet = workbook.create_sheet(sheet_name('ROOM', used))
    configure_sheet(sheet, panel_width, title, school_year)
    for row, value in [(1, institution), (2, college), (3, title),
                       (4, f'SY {school_year} • All programs • 90-minute display bands'),
                       (5, 'Shared bands favor the longest occupant; exact times apply. Red = actual overlap. TBA entries are listed separately.')]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=panel_width * 2)
        cell = write_value(sheet, row, 1, value)
        cell.font = Font(name='Arial Narrow', size=12 if row != 5 else 10, bold=row != 5)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        sheet.row_dimensions[row].height = wrapped_height([value], [panel_width * 2 * 25])

    print_areas = []
    panels = [('M', 'W'), ('T', 'Th'), ('F', 'Sa')]
    if 'Su' in days:
        panels.append(('Su',))
    for pair_index, pair in enumerate(panels):
        header = 6 + pair_index * (len(slots) + 2)
        for side, day in enumerate(pair):
            time_column = 1 + side * panel_width
            for column, value in enumerate(['Time'] + [f'{DAY_LABELS[day]}\n{room["name"]}' for room in rooms], time_column):
                cell = write_value(sheet, header, column, value)
                cell.font = Font(name='Arial Narrow', size=11, bold=True)
                cell.fill = PatternFill('solid', fgColor='DCE6F1')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
            sheet.row_dimensions[header].height = 36
            for row, (left, right) in enumerate(slots, header + 1):
                sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, 48)
                cell = write_value(sheet, row, time_column, f'{clock(left)}-{clock(right)}')
                cell.font = Font(name='Arial Narrow', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='F2F5F8')
                cell.border = border
                for column in range(time_column + 1, time_column + panel_width):
                    sheet.cell(row, column).border = border
            for column, room in enumerate(rooms, time_column + 1):
                offset = 0
                while offset < len(slots):
                    indexes = tuple(occupied[day, room['id'], offset])
                    if not indexes:
                        offset += 1
                        continue
                    until = offset + 1
                    overlap = conflicts[day, room['id'], offset]
                    while (until < len(slots) and tuple(occupied[day, room['id'], until]) == indexes
                           and conflicts[day, room['id'], until] == overlap):
                        until += 1
                    selected = [bookings[index] for index in indexes]
                    text = '\n\n'.join(label(entry, begin, finish) for entry, begin, finish, _ in selected)
                    if len(text) > 32767:
                        raise HTTPException(413, f'{room["name"]}: too many bookings in one time block to export safely.')
                    first, last = header + 1 + offset, header + until
                    cell = write_value(sheet, first, column, text)
                    cell.font = Font(name='Arial Narrow', size=10)
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.fill = PatternFill('solid', fgColor='FCE4D6' if overlap else 'DCE6F1')
                    cell.border = border
                    # Short or stacked classes need enough height to retain every label.
                    height = wrapped_height([text], [25]) / (until - offset)
                    if height > 409:
                        raise HTTPException(413, f'{room["name"]}: too many booking details in one band to fit in Excel.')
                    for row in range(first, last + 1):
                        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height, height)
                        sheet.cell(row, column).fill = copy(cell.fill)
                    if first != last:
                        sheet.merge_cells(start_row=first, start_column=column, end_row=last, end_column=column)
                    offset = until
            if not rooms:
                write_value(sheet, header, time_column + 1, DAY_LABELS[day])
                write_value(sheet, header + 1, time_column + 1, 'No rooms to export.')
            for chunk in range(0, max(1, len(rooms)), ROOMS_PER_PAGE):
                first_column = time_column + 1 + chunk
                last_column = time_column + min(chunk + ROOMS_PER_PAGE, panel_width - 1)
                print_areas.append(f'{get_column_letter(first_column)}{header}:'
                                   f'{get_column_letter(last_column)}{header + len(slots)}')
    sheet.print_area = print_areas
    if unplotted:
        sheet = workbook.create_sheet(sheet_name('Room TBA', used))
        headers = ['Course Code', 'Section', 'Program', 'Time', 'Days', 'Room', 'Faculty', 'Reason']
        widths = [20, 24, 24, 26, 18, 24, 30, 42]
        sheet.freeze_panes = 'A3'
        sheet.merge_cells('A1:H1')
        write_value(sheet, 1, 1, 'Entries not plotted: missing room, day, or valid time').font = Font(bold=True, size=12)
        for column, (header, width) in enumerate(zip(headers, widths), 1):
            cell = write_value(sheet, 2, column, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='DCE6F1')
            sheet.column_dimensions[cell.column_letter].width = width
        for row, entry in enumerate(unplotted, 3):
            reason = []
            if entry.get('room_id') not in room_ids:
                reason.append('Room not assigned or unavailable')
            if not interval(entry):
                reason.append('Time TBA or invalid')
            if not meeting_days(entry):
                reason.append('Days TBA or invalid')
            values = [entry.get('Course Code'), entry.get('Section'), entry.get('Program'),
                      entry.get('Time (24 Hrs)') or entry.get('Time (LPU Std)') or 'TBA',
                      entry.get('Days') or 'TBA', entry.get('Room') or 'TBA', entry.get('Faculty') or 'TBA', '; '.join(reason)]
            for column, value in enumerate(values, 1):
                cell = write_value(sheet, row, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                cell.border = border
            sheet.row_dimensions[row].height = wrapped_height(values, widths)
        sheet.auto_filter.ref = f'A2:H{sheet.max_row}'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = '1:2'
        sheet.print_area = f'A1:H{sheet.max_row}'


def build_workbook(rooms, entries, school_year, institution=DEFAULT_INSTITUTION,
                   college=DEFAULT_COLLEGE, title='ROOM SCHEDULE', include_sunday=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheets(workbook, rooms, entries, school_year, institution, college, title, include_sunday)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
